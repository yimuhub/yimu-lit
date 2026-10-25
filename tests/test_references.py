"""参考资料库测试（全离线）。

覆盖：格式抽取 / 切片偏移 / 体裁判定 / 入库全链路 / LLM 降级 / 检索 / 上下文注入 /
查重 / URL 抓取（monkeypatch）/ 并发 / 路径安全。

所有落盘都指向 tmp_path；LLM 由 conftest 统一 stub 为 {}，需要“正常 JSON”的用例
在本文件内再 monkeypatch 一次。
"""
import json
import os
import threading
import zipfile

import pytest

from bot import ingest, llm, references, writer
from bot.safeio import atomic_write_json


@pytest.fixture
def store(tmp_path, monkeypatch):
    refs = tmp_path / "references"
    projects = tmp_path / "projects"
    refs.mkdir()
    projects.mkdir()
    monkeypatch.setattr(references, "REFERENCES_DIR", str(refs))
    monkeypatch.setattr(writer, "PROJECTS_DIR", str(projects))
    references._clear_cache()
    yield {"refs": refs, "projects": projects}
    references._clear_cache()


def make_project(store, key="proj-1"):
    pdir = store["projects"] / key
    pdir.mkdir(parents=True, exist_ok=True)
    atomic_write_json(str(pdir / "02-写作计划.json"), {"config": {}, "meta": {}})
    return pdir


def write_bytes(tmp_path, name, data):
    p = tmp_path / name
    p.write_bytes(data)
    return str(p)


def make_min_pdf(path, text="Hello Reference PDF Document"):
    """手写最小合法 PDF（计算 xref 偏移），供 pypdf 抽取测试。"""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = ("BT /F1 18 Tf 20 100 Td (" + text + ") Tj ET").encode()
    objs[3] = (b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
               + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    with open(path, "wb") as f:
        f.write(out)


def make_min_epub(path):
    container = (
        '<?xml version="1.0"?><container version="1.0" '
        'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
        '<rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>')
    opf = (
        '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0">'
        '<manifest>'
        '<item id="a" href="texts/a.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="z" href="texts/z.xhtml" media-type="application/xhtml+xml"/>'
        '</manifest><spine><itemref idref="z"/><itemref idref="a"/></spine></package>')
    a = "<html><body><h1>SECOND</h1><p>Alpha chapter body.</p></body></html>"
    z = "<html><body><h1>FIRST</h1><p>Zeta chapter body.</p></body></html>"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("META-INF/container.xml", container)
        zf.writestr("OEBPS/content.opf", opf)
        zf.writestr("OEBPS/texts/a.xhtml", a)
        zf.writestr("OEBPS/texts/z.xhtml", z)


# ---------------------------------------------------------------- 1. 纯文本/HTML

def test_txt_encoding_and_clean(tmp_path):
    p = write_bytes(tmp_path, "a.txt", "第一段。\r\n\r\n\r\n第二段。  \n".encode())
    r = ingest.extract_text(p)
    assert r["ok"] and r["ext"] == ".txt"
    assert "第一段。" in r["text"] and "第二段。" in r["text"]
    assert "\r" not in r["text"] and "\n\n\n" not in r["text"]

    gb = write_bytes(tmp_path, "b.txt", "这是国标编码的中文资料，用于测试解码。".encode("gb18030"))
    r2 = ingest.extract_text(gb)
    assert r2["ok"] and "国标编码" in r2["text"]
    assert any("gb18030" in w for w in r2["warnings"])


def test_md_and_html_clean(tmp_path):
    md = write_bytes(tmp_path, "a.md", "# 标题\n\n正文第一段。\n\n正文第二段。".encode())
    r = ingest.extract_text(md)
    assert r["ok"] and r["kind"] in ("novel", "prose", "reference")
    assert "正文第二段" in r["text"]

    html_bytes = (
        "<html><head><style>.x{color:red}</style><script>alert(1)</script></head>"
        "<body><nav>首页 登录 注册</nav><h1>历史资料</h1>"
        "<p>贞观年间的盐铁税收制度记载于旧唐书。</p>"
        "<footer>版权所有 广告</footer></body></html>"
    ).encode()
    hp = write_bytes(tmp_path, "a.html", html_bytes)
    r2 = ingest.extract_text(hp)
    assert r2["ok"]
    assert "贞观年间" in r2["text"]
    assert "alert" not in r2["text"] and "color:red" not in r2["text"]
    assert "首页" not in r2["text"] and "广告" not in r2["text"]


def test_ext_reject_and_size_limit(tmp_path, monkeypatch):
    bad = write_bytes(tmp_path, "a.exe", b"MZ")
    r = ingest.extract_text(bad)
    assert not r["ok"] and "不支持" in r["error"]

    big = write_bytes(tmp_path, "b.txt", b"x" * 64)
    r2 = ingest.extract_text(big, max_bytes=16)
    assert not r2["ok"] and "上限" in r2["error"]

    # add_upload 走的是运行时常量，确保部署时调小上限也生效。
    monkeypatch.setattr(ingest, "MAX_BYTES", 8)
    r3 = references.add_upload(b"123456789", "big.txt")
    assert not r3["ok"] and "上限" in r3["error"]


# ---------------------------------------------------------------- 2. docx / pdf / epub

def test_docx_extract(tmp_path):
    docx = pytest.importorskip("docx")
    from docx import Document
    p = str(tmp_path / "a.docx")
    doc = Document()
    doc.add_heading("人物设定", level=1)
    doc.add_paragraph("主角林砚是一名验尸官，习惯在雨天记录尸格。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "林砚"
    doc.save(p)
    r = ingest.extract_text(p)
    assert r["ok"] and r["ext"] == ".docx"
    assert "验尸官" in r["text"] and "林砚" in r["text"]
    assert docx is not None


def test_pdf_extract_or_skip(tmp_path):
    pytest.importorskip("pypdf")
    p = str(tmp_path / "a.pdf")
    make_min_pdf(p)
    r = ingest.extract_text(p)
    assert r["ok"], r.get("error")
    assert "Hello Reference PDF" in r["text"]
    assert r.get("pages") == 1


def test_epub_spine_order(tmp_path):
    p = str(tmp_path / "a.epub")
    make_min_epub(p)
    r = ingest.extract_text(p)
    assert r["ok"], r.get("error")
    assert r["text"].index("FIRST") < r["text"].index("SECOND")
    assert r["text"].index("Zeta") < r["text"].index("Alpha")


# ---------------------------------------------------------------- 3. 切片

def test_chunk_offsets_overlap_and_hard_split():
    paras = [f"第{i}段。" + "内容文字" * 45 for i in range(1, 7)]
    text = "\n\n".join(paras)
    chunks = ingest.chunk_text(text, size=420, overlap=150)
    assert len(chunks) > 2
    for c in chunks:
        assert text[c["start"]:c["end"]] == c["text"]
        assert c["no"] >= 1
    assert chunks[1]["start"] < chunks[0]["end"]  # 段落边界重叠
    assert chunks[-1]["end"] == len(text)

    long_para = "甲" * 5000
    hard = ingest.chunk_text(long_para, size=1000, overlap=100, max_chunks=5000)
    assert len(hard) == 5
    assert all(c.get("hard_split") for c in hard[:-1])
    for c in hard:
        assert long_para[c["start"]:c["end"]] == c["text"]
    assert hard[-1]["end"] == len(long_para)

    # 空文本/参数边界不会死循环。
    assert ingest.chunk_text("", size=100) == []
    assert ingest.chunk_text("一段话", size=0) == []
    assert len(ingest.chunk_text(text, size=250, overlap=200, max_chunks=2)) == 2


# ---------------------------------------------------------------- 4. 体裁判定

def test_detect_kind():
    novel = ("第1章 雨夜\n\n“你来晚了。”他说。\n\n" + "门外的雨声压过了脚步。\n\n") * 8
    assert ingest.detect_kind(novel) == "novel"

    script = ("INT. 咖啡馆 - 日\n\n林砚：你迟到了。\n\n苏晚：路上堵车。\n\n"
              "EXT. 街道 - 夜\n\nCUT TO:\n\n林砚（画外音）：那天的雨很大。\n")
    assert ingest.detect_kind(script) == "script"

    poetry = "七律 秋思\n平仄平平仄仄平，\n押韵下平十一尤。\n风急天高猿啸哀，\n渚清沙白鸟飞回。"
    assert ingest.detect_kind(poetry) == "poetry"

    prose = "散文 故乡的河\n\n我常常想起故乡的那条河，河水在夏天涨起来，漫过石阶，带着青草的气味。"
    assert ingest.detect_kind(prose) == "prose"

    reference = ("历史资料：唐代盐铁专卖制度规定，盐商必须持有盐引才能运销，"
                 "朝廷在主要产区设置盐监，负责称量、课税与缉私。")
    assert ingest.detect_kind(reference) == "reference"


# ---------------------------------------------------------------- 5. 入库全链路

def test_add_list_get_chunks_usage_delete(store):
    make_project(store)
    r = references.add_text("青铜面具出土于祭祀坑，表面有饕餮纹。\n\n祭司在冬至举行祭祀仪式。",
                            "青铜考", usage_mode="facts")
    assert r["ok"], r.get("error")
    rid = r["ref_id"]
    assert references.REF_ID_RE.match(rid)
    meta_dir = store["refs"] / rid
    assert (meta_dir / "meta.json").is_file()
    assert (meta_dir / "text.txt").is_file()
    assert (meta_dir / "chunks.jsonl").is_file()
    assert (meta_dir / "analysis.json").is_file()

    listed = references.list_refs()
    assert listed["ok"] and listed["stats"]["count"] == 1
    assert listed["items"][0]["title"] == "青铜考"

    ref = references.get_ref(rid)
    assert ref and ref["id"] == rid and "text" not in ref
    preview = references.read_chunks(rid, 0, 10)
    assert preview["ok"] and preview["items"][0]["text"]
    assert "青铜面具" in preview["items"][0]["text"]

    link = references.link_to_project("proj-1", rid)
    assert link["ok"] and references.project_ref_ids("proj-1") == [rid]
    assert references.list_refs("proj-1")["items"][0]["linked"] is True

    upd = references.set_usage(rid, "adapt")
    assert upd["ok"] and references.get_ref(rid)["usage_mode"] == "adapt"
    assert references.project_refs("proj-1")[0]["usage_mode"] == "adapt"
    assert not references.set_usage(rid, "bogus")["ok"]

    deleted = references.delete_ref(rid)
    assert deleted["ok"] and deleted["deleted"] is True
    assert references.project_ref_ids("proj-1") == []
    assert references.list_refs()["stats"]["count"] == 0
    again = references.delete_ref(rid)
    assert again["ok"] and again["deleted"] is False


def test_add_upload_and_magic(store):
    r = references.add_upload("上传的中文资料。".encode(), "../../危险.md")
    assert r["ok"], r.get("error")
    rid = r["ref_id"]
    files = os.listdir(store["refs"] / rid)
    assert any(f.startswith("source.") for f in files)
    assert references.get_ref(rid)["source_type"] == "upload"

    bad = references.add_upload(b"not a real pdf", "x.pdf")
    assert not bad["ok"] and "文件头" in bad["error"]
    bad_ext = references.add_upload(b"abc", "x.exe")
    assert not bad_ext["ok"] and "不支持" in bad_ext["error"]


# ---------------------------------------------------------------- 6. LLM 分析降级

def test_analyze_fallback_and_success(store, monkeypatch):
    r = references.add_text("唐代盐铁专卖制度规定，盐商必须持有盐引才能运销。", "盐铁制度")
    rid = r["ref_id"]

    fb = references.analyze_ref(rid)
    assert fb["ok"] and fb["fallback"] is True
    assert "盐铁制度" in fb["analysis"]["summary"]
    ref = references.get_ref(rid)
    assert ref["analysis"]["fallback"] is True
    assert ref["analysis_at"]

    good = {
        "summary": "介绍唐代盐铁专卖与盐引制度。",
        "genre": "历史资料",
        "structure": "先总述制度，再说明盐引流程。",
        "characters": ["盐商", "盐监"],
        "style_traits": ["条目化、术语密集"],
        "key_facts": ["盐商必须持盐引运销"],
        "caution": ["需核对史料版本"],
    }
    monkeypatch.setattr(llm, "chat_json", lambda *a, **k: good)
    ok = references.analyze_ref(rid)
    assert ok["ok"] and ok["fallback"] is False
    ref = references.get_ref(rid)
    assert ref["analysis"]["genre"] == "历史资料"
    assert "盐引流程" in ref["analysis"]["structure"]
    assert ref["summary"] == "介绍唐代盐铁专卖与盐引制度。"
    raw = json.loads((store["refs"] / rid / "analysis.json").read_text(encoding="utf-8"))
    assert raw["analysis"]["style_traits"] == ["条目化、术语密集"]


# ---------------------------------------------------------------- 7. 检索

def test_retrieve_bm25_filters(store):
    a = references.add_text("唐代盐铁专卖制度规定，盐商必须持有盐引才能运销，盐监负责课税。",
                            "盐铁")
    b = references.add_text("青铜面具出土于三星堆祭祀坑，表面有饕餮纹与云雷纹，用于祭祀。",
                            "青铜")
    c = references.add_text("宋代瓷器以汝窑青瓷最为珍贵，釉色天青，开片细密如蟹爪。",
                            "汝窑")
    assert a["ok"] and b["ok"] and c["ok"]

    hits = references.retrieve("青铜面具 饕餮纹", k=3)
    assert hits and hits[0]["ref_id"] == b["ref_id"]
    assert hits[0]["usage_mode"] == "facts"
    assert "饕餮" in hits[0]["snippet"]

    hits2 = references.retrieve("汝窑青瓷 釉色", k=3)
    assert hits2 and hits2[0]["ref_id"] == c["ref_id"]

    # disabled 永不返回
    references.set_usage(b["ref_id"], "disabled")
    hits3 = references.retrieve("青铜面具 饕餮纹", k=5)
    assert all(h["ref_id"] != b["ref_id"] for h in hits3)

    # ref_ids 限定生效：这两篇里没有“饕餮纹”
    hits4 = references.retrieve("饕餮纹", ref_ids=[a["ref_id"], c["ref_id"]], k=5)
    assert hits4 == []


# ---------------------------------------------------------------- 8. 上下文注入

def test_context_block_rules_limits_and_style(store):
    text = ("雨夜里的青石巷深处传来铜铃声响，老掌柜把最后一盏油灯移到柜台上，"
            "低声嘱咐学徒守住后门的钥匙与账册，谁也别放进。")
    r = references.add_text(text, "巷口旧事")
    rid = r["ref_id"]

    block = references.context_block("青石巷 铜铃 老掌柜", [rid], max_chars=800, k=3)
    assert "[参考：巷口旧事#" in block
    assert "不得连续复制" in block
    assert len(block) <= 800

    tight = references.context_block("青石巷 铜铃", [rid], max_chars=120, k=3)
    assert len(tight) <= 120

    assert references.set_usage(rid, "style")["ok"]
    style_block = references.context_block("青石巷 铜铃 老掌柜", [rid], max_chars=1200, k=3)
    assert "用法：style" in style_block
    assert text[:40] not in style_block  # style 不给原文情节
    assert text[10:50] not in style_block

    assert references.context_block("青石巷", [], max_chars=800) == ""


# ---------------------------------------------------------------- 9. 查重

def test_overlap_check_copy_and_original(store):
    passage = ("雨夜里的青石巷深处传来铜铃声响，老掌柜把最后一盏油灯移到柜台上，"
               "低声嘱咐学徒守住后门的钥匙与账册，谁也别放进。")
    r = references.add_text(passage, "旧事")
    assert r["ok"]

    generated = "故事开始。" + passage + "故事到这里结束。"
    res = references.overlap_check(generated, n=12)
    assert res["match_count"] >= 1
    assert res["ratio"] > 0
    assert res["ok"] is False
    m = res["matches"][0]
    assert m["ref_id"] == r["ref_id"] and m["title"] == "旧事"
    assert m["chunk_no"] >= 1 and m["excerpt"] and m["length"] >= 12
    assert set(res) >= {"ok", "ratio", "match_count", "matches", "note"}

    original = ("清晨的河岸上，渔夫收起湿漉漉的渔网，远处传来汽笛声。"
                "他抬头看了看云层，决定今天不出海。")
    clean = references.overlap_check(original, n=12)
    assert clean["match_count"] == 0 and clean["ratio"] == 0 and clean["ok"] is True


# ---------------------------------------------------------------- 10. URL 抓取

def test_add_url_success_empty_and_ssrf(store, monkeypatch):
    url = "http://93.184.216.34/reference"
    monkeypatch.setattr(references.distill, "fetch_url_text",
                        lambda u, max_chars=12000: "网页正文：唐代盐引制度与盐商运销流程。" * 3)
    r = references.add_url(url, "网页盐法")
    assert r["ok"], r.get("error")
    ref = references.get_ref(r["ref_id"])
    assert ref["source_type"] == "url" and ref["source"] == url

    monkeypatch.setattr(references.distill, "fetch_url_text",
                        lambda u, max_chars=12000: "   ")
    empty = references.add_url("http://93.184.216.34/empty")
    assert not empty["ok"] and "抓取失败" in empty["error"]

    called = []

    def boom(*a, **k):
        called.append(a)
        raise AssertionError("SSRF 地址不应触发抓取")

    monkeypatch.setattr(references.distill, "fetch_url_text", boom)
    for bad in ("http://127.0.0.1/x", "http://[::1]/x", "file:///etc/passwd"):
        out = references.add_url(bad)
        assert not out["ok"] and out["error"]
    assert called == []


# ---------------------------------------------------------------- 11. 并发新增

def test_concurrent_add_text(store):
    results = []
    lock = threading.Lock()

    def work(i):
        out = references.add_text(f"并发测试文本第{i}号，内容唯一标记 marker{i}。" * 20,
                                  f"并发{i}")
        with lock:
            results.append(out)

    threads = [threading.Thread(target=work, args=(i,)) for i in range(8)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert len(results) == 8 and all(r["ok"] for r in results)
    ids = {r["ref_id"] for r in results}
    assert len(ids) == 8
    assert references.list_refs()["stats"]["count"] == 8
    for rid in ids:
        meta = json.loads((store["refs"] / rid / "meta.json").read_text(encoding="utf-8"))
        assert meta["id"] == rid and meta["chunks"] >= 1
    leftovers = [p for p in store["refs"].rglob("*") if p.name.startswith(".tmp-")]
    assert leftovers == []


# ---------------------------------------------------------------- 12. 路径安全

def test_path_traversal_blocked(store):
    assert references.get_ref("../../etc/passwd") is None
    assert references.get_ref("r-../../x") is None
    assert references.get_ref("r-0123456789ab/../x") is None
    out = references.delete_ref("r-../../x")
    assert not out["ok"] and out["error"]
    out2 = references.delete_ref("../../etc/passwd")
    assert not out2["ok"] and out2["error"]
    assert references.read_chunks("../x")["ok"] is False
    assert references.set_usage("r-../../x", "facts")["ok"] is False

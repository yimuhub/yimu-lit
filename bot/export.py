"""专业导出：Markdown / DOCX / EPUB / 打印 HTML，剧本另出 Fountain / Final Draft(.fdx)。

只依赖标准库；DOCX 用 python-docx（可选依赖，缺失时给出清晰错误）。
"""
from __future__ import annotations

import html
import io
import os
import re
import uuid
import zipfile

from . import writer as novel_writer

_SCENE_RE = re.compile(r"^(内景|外景|内/外景|日景|夜景|INT\.?|EXT\.?|I/E)", re.I)
_NONAME_RE = re.compile(r"[，。！？,.!?：:；;、\s]")


# ---------------------------------------------------------------- 取文

def _strip_md(s: str) -> str:
    s = re.sub(r"[*_`]+", "", s or "")
    s = re.sub(r"^#+\s*", "", s)
    return s.strip()


def novel_chapters(proj) -> list[dict]:
    plan = proj.plan() or {}
    out = []
    for ch in plan.get("chapters") or []:
        fp = os.path.join(proj.path, os.path.basename(ch.get("filePath") or ""))
        try:
            text = open(fp, encoding="utf-8").read()
        except OSError:
            continue
        text = novel_writer._clean_chapter_text(text)
        if text.strip():
            out.append({"n": ch.get("chapterNumber"), "title": ch.get("title") or "",
                        "text": text})
    return out


def is_script(proj) -> bool:
    return os.path.exists(os.path.join(proj.path, "剧本计划.json"))


def script_scenes(proj) -> list[dict]:
    plan = proj.plan() or {}
    out = []
    for sc in plan.get("scenes") or []:
        fn = os.path.basename(sc.get("fileName") or "")
        fp = os.path.join(proj.body_dir(), fn) if fn else ""
        if not fp:
            continue
        try:
            text = open(fp, encoding="utf-8").read().strip()
        except OSError:
            continue
        if text:
            out.append({"no": sc.get("no"), "heading": sc.get("heading") or "", "text": text})
    return out


# ---------------------------------------------------------------- Markdown

def markdown(proj) -> str:
    if is_script(proj):
        plan = proj.plan() or {}
        name = plan.get("scriptName") or proj.name
        parts = [f"# {name}\n\n> {plan.get('form', '')} · {len(plan.get('scenes') or [])} 场"]
        for sc in script_scenes(proj):
            parts.append(f"## 第{sc['no']}场 {sc['heading']}\n\n{sc['text']}")
        return "\n\n---\n\n".join(parts)
    plan = proj.plan() or {}
    name = plan.get("novelName") or proj.name
    lines = [f"# {name}", f"> 共 {len(plan.get('chapters') or [])} 章"]
    for ch in novel_chapters(proj):
        lines.append(f"## 第{ch['n']}章 {_strip_md(ch['title'])}\n\n{ch['text']}")
    return "\n\n---\n\n".join(lines)


# ---------------------------------------------------------------- DOCX

def _docx_para(doc, text: str):
    from docx.shared import Pt
    p = doc.add_paragraph()
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = p.add_run(part[2:-2] if bold else _strip_md(part))
        run.bold = bold
        run.font.size = Pt(12)
    return p


def docx_bytes(title: str, subtitle: str, blocks: list[tuple[int, str]]) -> bytes:
    """blocks: [(级别, 文本)]，级别 1=章标题，2=小节，0=正文/引用。"""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt
    except ImportError as e:  # pragma: no cover - 依赖缺失时的明确提示
        raise RuntimeError("导出 DOCX 需要 python-docx：pip install python-docx") from e

    doc = Document()
    for name in ("Normal", "Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = "宋体"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.styles["Normal"].font.size = Pt(11)

    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for level, text in blocks:
        if level == 1:
            doc.add_heading(_strip_md(text), level=1)
        elif level == 2:
            doc.add_heading(_strip_md(text), level=2)
        elif level == 3:
            doc.add_heading(_strip_md(text), level=3)
        else:
            for para in re.split(r"\n\s*\n", text):
                para = para.strip()
                if para and para != "---":
                    _docx_para(doc, para.replace("\n", " "))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _md_blocks(md_text: str) -> list[tuple[int, str]]:
    """把 Markdown 粗分成 (级别, 文本) 块，足够 DOCX/EPUB 使用。"""
    blocks, buf = [], []
    for line in md_text.splitlines():
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            if buf:
                blocks.append((0, "\n".join(buf)))
                buf = []
            blocks.append((len(m.group(1)), m.group(2)))
        elif line.strip() == "---":
            continue
        else:
            buf.append(re.sub(r"^>\s?", "", line))
    if buf:
        blocks.append((0, "\n".join(buf)))
    return blocks


# ---------------------------------------------------------------- XHTML

def _md_to_xhtml(text: str) -> str:
    out = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para or para == "---":
            continue
        safe = html.escape(para)
        safe = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", safe)
        safe = re.sub(r"^#{1,3}\s*", "", safe)
        out.append(f"<p>{safe.replace(chr(10), '<br/>')}</p>")
    return "\n".join(out)


_EPUB_CSS = """body{font-family:serif;line-height:1.8;margin:5%;}
h1,h2{text-align:center;margin:1.4em 0 1em;}
p{text-indent:2em;margin:0 0 .8em;}
strong{font-weight:bold;}"""


def epub_bytes(title: str, author: str, chapters: list[dict]) -> bytes:
    """EPUB 3（含 nav 目录）；章节：{n,title,text}。"""
    uid = "urn:uuid:" + str(uuid.uuid5(uuid.NAMESPACE_URL, title + "|" + author))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" '
                   'media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr("OEBPS/style.css", _EPUB_CSS)
        items, spine, nav = [], [], []
        for i, ch in enumerate(chapters, 1):
            fn = f"ch{i:03d}.xhtml"
            body = _md_to_xhtml(ch["text"])
            z.writestr(f"OEBPS/{fn}",
                       '<?xml version="1.0" encoding="UTF-8"?>'
                       '<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" '
                       'xml:lang="zh-CN" lang="zh-CN"><head><meta charset="utf-8"/>'
                       f'<title>第{ch["n"]}章</title><link rel="stylesheet" href="style.css"/>'
                       f'</head><body><h2>第{ch["n"]}章 {html.escape(_strip_md(ch["title"]))}</h2>'
                       f'<div>{body}</div></body></html>')
            items.append(f'<item id="c{i}" href="{fn}" media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="c{i}"/>')
            nav.append(f'<li><a href="{fn}">第{ch["n"]}章 '
                       f'{html.escape(_strip_md(ch["title"]))}</a></li>')
        z.writestr("OEBPS/nav.xhtml",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" '
                   'xmlns:epub="http://www.idpf.org/2007/ops" xml:lang="zh-CN" lang="zh-CN">'
                   '<head><meta charset="utf-8"/><title>目录</title>'
                   '<link rel="stylesheet" href="style.css"/></head><body>'
                   f'<h1>{html.escape(title)}</h1><nav epub:type="toc" id="toc"><h2>目录</h2><ol>'
                   + "".join(nav) + "</ol></nav></body></html>")
        z.writestr("OEBPS/content.opf",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid">'
                   '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
                   f'<dc:identifier id="uid">{uid}</dc:identifier>'
                   f'<dc:title>{html.escape(title)}</dc:title>'
                   f'<dc:creator>{html.escape(author)}</dc:creator>'
                   '<dc:language>zh-CN</dc:language>'
                   '<meta property="dcterms:modified">2026-01-01T00:00:00Z</meta></metadata>'
                   '<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" '
                   'properties="nav"/><item id="css" href="style.css" media-type="text/css"/>'
                   + "".join(items) +
                   f'</manifest><spine toc="nav">{"".join(spine)}</spine></package>')
    return buf.getvalue()


def print_html(title: str, author: str, chapters: list[dict]) -> str:
    body = []
    for ch in chapters:
        body.append(f'<section class="ch"><h2>第{ch["n"]}章 '
                    f'{html.escape(_strip_md(ch["title"]))}</h2>{_md_to_xhtml(ch["text"])}</section>')
    return """<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"/>
<title>%s</title><style>
body{max-width:42em;margin:0 auto;padding:2em 1.4em;font-family:"Songti SC","Noto Serif CJK SC",serif;
line-height:1.9;color:#111;background:#fff;}
h1{text-align:center;font-size:2em;margin:1em 0 .2em;}
h2{margin:2em 0 1em;font-size:1.4em;page-break-after:avoid;}
p{text-indent:2em;margin:0 0 .85em;text-align:justify;}
.author{text-align:center;color:#555;margin-bottom:2em;}
.ch{page-break-before:always;}
@media print{@page{margin:2cm;}body{padding:0;}}
</style></head><body><h1>%s</h1><div class="author">%s</div>%s</body></html>""" % (
        html.escape(title), html.escape(title), html.escape(author), "".join(body))


# ---------------------------------------------------------------- 剧本 Fountain / FDX

def _classify(lines: list[str]):
    for i, line in enumerate(lines):
        t = line.strip()
        if not t:
            yield "blank", ""
            continue
        prev = lines[i - 1].strip() if i else ""
        nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
        if i == 0 and _SCENE_RE.match(t):
            yield "scene", t
        elif re.fullmatch(r"[（(].{0,40}[）)]", t):
            yield "paren", t[1:-1]
        elif 2 <= len(t) <= 10 and not _NONAME_RE.search(t) and nxt:
            yield "char", t
        elif prev and (len(prev.strip()) <= 10 and not _NONAME_RE.search(prev.strip())):
            yield "dialogue", t
        else:
            yield "action", t


def fountain(proj) -> str:
    out = []
    for sc in script_scenes(proj):
        lines = sc["text"].splitlines()
        for kind, text in _classify(lines):
            if kind == "blank":
                out.append("")
            elif kind == "scene":
                out.append("." + text.upper())
                out.append("")
            elif kind == "paren":
                out.append("(" + text + ")")
            elif kind == "char":
                out.append("@" + text)
            else:
                out.append(text)
        out.append("")
    return "\n".join(out).strip() + "\n"


def fdx_bytes(proj) -> bytes:
    title = (proj.plan() or {}).get("scriptName") or proj.name
    ps = []
    for sc in script_scenes(proj):
        for kind, text in _classify(sc["text"].splitlines()):
            if kind == "blank" or not text:
                continue
            ptype = {"scene": "Scene Heading", "action": "Action", "char": "Character",
                     "paren": "Parenthetical", "dialogue": "Dialogue"}[kind]
            ps.append(f'<Paragraph Type="{ptype}"><Text>{html.escape(text)}</Text></Paragraph>')
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
           '<FinalDraft DocumentType="Script" Template="No" Version="1">\n'
           f'<Title>{html.escape(title)}</Title>\n<Content>\n' + "\n".join(ps) +
           "\n</Content>\n</FinalDraft>\n")
    return xml.encode("utf-8")
